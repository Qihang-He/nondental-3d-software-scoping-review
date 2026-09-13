# -*- coding: utf-8 -*-
# --- path bootstrap added by the repository build -------------------------
# Resolves the project root from the SCOPING_ROOT environment variable, or from
# this file's location when the script sits three levels below the project root.
import os as _os
_ROOTP = (_os.environ.get('SCOPING_ROOT')
          or _os.path.dirname(_os.path.dirname(_os.path.dirname(
              _os.path.abspath(__file__)))))
# -------------------------------------------------------------------------
"""
_tools_supp2_v2.py —— 重建 Supplementary File 2（AI 辅助筛选与编码的完整可核验输出）
工作表：
  1_Screening_exclusion_reasons  1,639 条筛查阶段排除记录及其单一原因
  2_Run_consistency              2,556 条三轮复筛标签与多数标签
  3_Outcome_coding               863 条纳入研究的优势/挑战/缺口编码
  4_Codebook                     全部代码与操作性定义
  5_Model_and_run_metadata       模型、解码参数、提示词哈希、运行时间
输出：02_图表附件/R2_补充材料/Supplementary_File_2_R2.xlsx
"""
import os
import re
import json
import hashlib
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = _ROOTP
ANA = os.path.join(ROOT, _os.path.join(_ROOTP, '03_数据'), '08_分析用')
OUTD = os.path.join(ROOT, _os.path.join(_ROOTP, '02_图表附件'), 'R2_补充材料')
os.makedirs(OUTD, exist_ok=True)

REASON_LBL = {1: 'Outside the scope of dentistry',
              2: 'Not an original research report (review, editorial, conference abstract)',
              3: 'No non-dental 3D software identified',
              4: 'Outside the prespecified date window',
              5: 'Other / not classifiable'}

ADV = {'A1': 'Greater geometric freedom or customisation',
       'A2': 'Analysis functions unavailable in dental platforms',
       'A3': 'Cost savings, free or open-source access',
       'A4': 'Improved interoperability or data exchange',
       'A5': 'Automation and time efficiency',
       'A6': 'Improved accuracy or measurement precision',
       'A9': 'No advantage explicitly reported'}
CHA = {'C1': 'Steep learning curve or training requirement',
       'C2': 'File-format or interoperability problems',
       'C3': 'Licence cost or access restrictions',
       'C4': 'Limited validation evidence',
       'C5': 'Time-consuming workflow',
       'C6': 'Software instability, bugs or errors',
       'C7': 'High hardware or computational demand',
       'C8': 'Regulatory, approval or medico-legal uncertainty',
       'C9': 'No challenge explicitly reported'}
GAP = {'G1': 'Need for clinical validation',
       'G2': 'Need for automation or AI integration',
       'G3': 'Need for standardisation and reporting guidance',
       'G4': 'Need for larger or multi-centre samples',
       'G9': 'No gap explicitly reported'}

# ---------- 1 排除原因 ----------
rz = pd.read_csv(os.path.join(ANA, 'prisma_pass', 'reason_parsed.csv'), low_memory=False)
rz = rz.drop_duplicates('_nt')
FIN = pd.read_csv(os.path.join(ANA, '分析数据集_final_v4.csv'), low_memory=False)
inc_keys = set(FIN['Key'].astype(str))
added = set()
for _cand in ('最终新增纳入.csv', '新增纳入_逐条溯源.csv', '最终新增纳入_863.csv'):
    _p = os.path.join(ANA, _cand)
    if os.path.exists(_p):
        _t = pd.read_csv(_p, low_memory=False)
        if 'Key' in _t.columns:
            added |= set(_t['Key'].astype(str))
rz['_reassessed'] = rz['Key'].astype(str).isin(added)
if rz['_reassessed'].sum() not in (301, 304):
    print('⚠ 复核集合命中 %d 条（期望 304 / 保留 301）' % rz['_reassessed'].sum())
s1 = rz.copy()
s1['Subsequently re-assessed in full text and re-classified as eligible'] = \
    s1['_reassessed'].map({True: 'Yes', False: 'No'})
winset = set()
_wf = os.path.join(ANA, '时间窗外补充剔除_v4.csv')
if os.path.exists(_wf):
    winset = set(pd.read_csv(_wf)['Title'].astype(str))


def _status(r):
    if not r['_reassessed']:
        return 'Excluded at title/abstract screening'
    if str(r['Key']) in inc_keys:
        return 'Included in the review after full-text re-assessment'
    return 'Excluded after full-text re-assessment (outside the date window)'


s1['Final status'] = s1.apply(_status, axis=1)
s1 = pd.DataFrame({
    'Record (title)': s1['Title'],
    'Assigned primary exclusion reason (code)': s1['原因代码'],
    'Assigned primary exclusion reason (label)': s1['原因代码'].map(REASON_LBL),
    'Subsequently re-assessed in full text and re-classified as eligible':
        s1['Subsequently re-assessed in full text and re-classified as eligible'],
    'Final status': s1['Final status'],
}).reset_index(drop=True)
print('sheet1 screening exclusions (all title/abstract exclusions): %d' % len(s1))
print(s1['Final status'].value_counts().to_string())

# ---------- 2 三轮一致性 ----------
lc = pd.read_csv(os.path.join(ROOT, _os.path.join(_ROOTP, '07_AI重跑原始记录'), '一致性分析',
                              '三轮逐条结果与多数标签_v2.csv'), low_memory=False)
s2 = pd.DataFrame({'Record (title)': lc['Title'], 'Run 1': lc['run1'], 'Run 2': lc['run2'],
                   'Run 3': lc['run3'], 'Unanimous': lc['一致'],
                   'Majority label': lc['多数标签']})

# ---------- 3 结局编码 ----------
rq = pd.read_csv(os.path.join(ANA, 'rq3_pass', 'rq3_parsed.csv'), low_memory=False)
v3 = pd.read_csv(os.path.join(ANA, '分析数据集_final_v4.csv'), low_memory=False)
rq['Key'] = rq['Key'].astype(str)
v3['Key'] = v3['Key'].astype(str)
rq = rq[rq['Key'].isin(set(v3['Key']))].drop_duplicates('Key')
m = v3.set_index('Key')[['Journal', 'Year', 'study_type', 'Software Used']]
rq = rq.join(m, on='Key')
s3 = pd.DataFrame({
    'Record (title)': rq['Title'],
    'Journal': rq['Journal'],
    'Year': rq['Year'],
    'Study design': rq['study_type'],
    'Software used': rq['Software Used'],
    'Advantage codes': rq['优势'],
    'Advantage labels': rq['优势'].map(
        lambda v: '; '.join(ADV.get(x.strip().upper(), x) for x in str(v).split(',') if x.strip())),
    'Challenge codes': rq['挑战'],
    'Challenge labels': rq['挑战'].map(
        lambda v: '; '.join(CHA.get(x.strip().upper(), x) for x in str(v).split(',') if x.strip())),
    'Gap codes': rq['缺口'],
    'Gap labels': rq['缺口'].map(
        lambda v: '; '.join(GAP.get(x.strip().upper(), x) for x in str(v).split(',') if x.strip())),
})

# ---------- 4 代码本 ----------
rows = []
for grp, d in [('Advantage', ADV), ('Challenge', CHA), ('Gap', GAP)]:
    for k, v in d.items():
        rows.append({'Category': grp, 'Code': k, 'Definition': v})
rows += [
    {'Category': 'Exclusion reason', 'Code': str(k), 'Definition': v}
    for k, v in REASON_LBL.items()]
rows += [
    {'Category': 'Screening label', 'Code': 'include',
     'Definition': 'Meets all eligibility criteria with no exclusion factor apparent'},
    {'Category': 'Screening label', 'Code': 'exclude',
     'Definition': 'Clearly fails an eligibility criterion or meets an exclusion criterion'},
    {'Category': 'Screening label', 'Code': 'unsure',
     'Definition': 'Insufficient information in title and abstract to decide'},
    {'Category': 'Study design', 'Code': 'computational',
     'Definition': 'Computational, simulation or software-development study'},
    {'Category': 'Study design', 'Code': 'in_vitro',
     'Definition': 'In vitro or laboratory study'},
    {'Category': 'Study design', 'Code': 'clinical', 'Definition': 'Study involving patients'},
    {'Category': 'Study design', 'Code': 'technical_note',
     'Definition': 'Technical or dental technique report'},
    {'Category': 'Study design', 'Code': 'case_report', 'Definition': 'Case report or case series'},
    {'Category': 'Study design', 'Code': 'educational',
     'Definition': 'Educational or training study'},
]
s4 = pd.DataFrame(rows)

# ---------- 5 模型与运行元数据 ----------
prompt_fp = os.path.join(ROOT, _os.path.join(_ROOTP, '07_AI重跑原始记录'), 'system_prompt.txt')
ph = hashlib.sha256(open(prompt_fp, 'rb').read()).hexdigest()[:16]
S = json.load(open(os.path.join(ANA, '统计核心.json'), encoding='utf-8'))
s5 = pd.DataFrame({
    'Item': ['API endpoint', 'Requested model identifier', 'Model returned by the API',
             'Temperature', 'max_tokens', 'Streaming', 'Prompt SHA-256 (first 16 hex)',
             'Screening corpus', 'Screening runs retained', 'Run-to-run agreement',
             'Fleiss kappa (three runs)', 'Screening dates',
             'Exclusion-reason pass dates', 'Outcome-coding pass dates',
             'Raw response logging', 'Deterministic corpus definition', 'Sampling seed'],
    'Value': ['https://api.deepseek.com/v1/chat/completions',
              'deepseek-chat (rolling alias, not pinned to a version)',
              'deepseek-flash (DeepSeek-V4.1-Flash), as returned in every response',
              '0.1', '500', 'Disabled (stream = false)', ph,
              '2,556 unique records after duplicate removal',
              '3 independent runs with identical prompt and parameters',
              '94.3% unanimous; pairwise 95.1-98.2%',
              '0.936',
              '10 September 2026',
              '12 September 2026',
              '12 September 2026',
              'Request timestamp, returned model, response id, content and latency stored per record',
              'Unique normalised titles derived from the 2,572 Zotero records after de-duplication',
              '20260912'],
})

xlsx = os.path.join(OUTD, 'Supplementary_File_2_R2.xlsx')
with pd.ExcelWriter(xlsx, engine='openpyxl') as w:
    s1.to_excel(w, sheet_name='1_Screening_exclusion_reasons', index=False)
    s2.to_excel(w, sheet_name='2_Run_consistency', index=False)
    s3.to_excel(w, sheet_name='3_Outcome_coding', index=False)
    s4.to_excel(w, sheet_name='4_Codebook', index=False)
    s5.to_excel(w, sheet_name='5_Model_and_run_metadata', index=False)
    wb = w.book
    W = {'Record (title)': 70, 'Definition': 78, 'Item': 34, 'Value': 70,
         'Advantage labels': 46, 'Challenge labels': 46, 'Gap labels': 46,
         'Assigned primary exclusion reason (label)': 58}
    for sn in wb.sheetnames:
        ws = wb[sn]
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 30
        for i in range(1, ws.max_column + 1):
            h = ws.cell(row=1, column=i).value
            ws.column_dimensions[get_column_letter(i)].width = W.get(h, 20)
print('已生成:', xlsx)
print('sheets:', [x.shape for x in (s1, s2, s3, s4, s5)])
