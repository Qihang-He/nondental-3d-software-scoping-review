# -*- coding: utf-8 -*-
"""
_tools_md2docx.py —— Markdown -> Word（期刊投稿格式）

格式（Journal of Dentistry / Elsevier 通用要求）：
  * 单栏，Times New Roman 12 pt
  * 双倍行距；段前段后 0；无首行缩进（块状段落）
  * 连续行号
  * 关键：按「空行」分段，段内硬换行合并为同一段（避免每个换行都成为独立段落）
  * 标题（#/##/###）、无序/有序列表、表格（| ... |）、加粗 **...**

用法：python _tools_md2docx.py <in.md> <out.docx>
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Times New Roman'
SIZE = Pt(12)
LINE_SPACING = 2.0          # 双倍行距


def set_base_style(doc):
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = SIZE
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), FONT)
    rf.set(qn('w:hAnsi'), FONT)
    rf.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)


def add_line_numbers(doc):
    sectPr = doc.sections[0]._sectPr
    ln = OxmlElement('w:lnNumType')
    ln.set(qn('w:countBy'), '1')
    ln.set(qn('w:restart'), 'continuous')
    ln.set(qn('w:distance'), '360')
    pgmar = sectPr.find(qn('w:pgMar'))
    if pgmar is not None:
        pgmar.addnext(ln)
    else:
        sectPr.append(ln)


def add_runs(par, text):
    """处理 **粗体**、*斜体* 与 ^上标^。"""
    for seg in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\^[^\^]+\^)', text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**'):
            r = par.add_run(seg[2:-2])
            r.bold = True
        elif seg.startswith('^') and seg.endswith('^') and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            r.font.superscript = True
        elif seg.startswith('*') and seg.endswith('*') and len(seg) > 2:
            r = par.add_run(seg[1:-1])
            r.italic = True
        else:
            par.add_run(seg)


def fmt(par):
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.line_spacing = LINE_SPACING
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)
    return par


def heading(doc, text, level):
    h = doc.add_heading('', level=level)
    add_runs(h, text)
    for r in h.runs:
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.size = Pt({1: 14, 2: 12, 3: 12}.get(level, 12))
    fmt(h)
    return h


def md_to_docx(md_path, docx_path):
    lines = open(md_path, encoding='utf-8').read().replace('\r\n', '\n').split('\n')
    doc = Document()
    set_base_style(doc)

    buf = []

    def flush():
        if not buf:
            return
        text = ' '.join(buf).strip()
        buf.clear()
        if not text:
            return
        p = doc.add_paragraph()
        add_runs(p, text)
        fmt(p)

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and \
                re.match(r'^\|[\s:\-|]+\|$', lines[i + 1].strip()):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not re.match(r'^\|[\s:\-|]+\|$', lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        val = row[ci] if ci < len(row) else ''
                        cell = t.cell(ri, ci)
                        cell.text = ''
                        pr = cell.paragraphs[0]
                        add_runs(pr, val)
                        fmt(pr)
                        for run in pr.runs:
                            run.font.size = Pt(10)
                            if ri == 0:
                                run.bold = True
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            flush()
            heading(doc, m.group(2).strip(), min(len(m.group(1)), 3))
            i += 1
            continue

        # 水平线
        if re.match(r'^-{3,}$', stripped):
            i += 1
            continue

        # 列表
        if re.match(r'^[-*]\s+', stripped):
            flush()
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^[-*]\s+', '', stripped))
            fmt(p)
            i += 1
            continue
        if re.match(r'^\d+\.\s+', stripped):
            flush()
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s+', '', stripped))
            fmt(p)
            i += 1
            continue
        if stripped.startswith('>'):
            flush()
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip('>').strip())
            fmt(p)
            i += 1
            continue

        # 空行 -> 段落分隔
        if stripped == '':
            flush()
            i += 1
            continue

        # 参考文献行：[1] ... 自成一段
        if re.match(r'^\[\d+\]', stripped):
            flush()
            p = doc.add_paragraph()
            add_runs(p, stripped)
            fmt(p)
            i += 1
            continue

        # 短标签行（如 Funding: / Authors' contributions: / Data availability:）自成一段
        if re.match(r'^[A-Z][A-Za-z\'\-\. \(\)]{0,95}:$', stripped):
            flush()
            p = doc.add_paragraph()
            add_runs(p, stripped)
            fmt(p)
            i += 1
            continue

        # 带内容的短标签项（如 Qihang He: Conceptualization, ...）自成一段
        m3 = re.match(r'^([A-Z][A-Za-z\'\-\. ]{0,24}):\s+\S', stripped)
        if m3 and len(m3.group(1).split()) <= 3:
            flush()
            p = doc.add_paragraph()
            add_runs(p, stripped)
            fmt(p)
            i += 1
            continue

        buf.append(stripped)
        i += 1

    flush()
    add_line_numbers(doc)
    doc.save(docx_path)
    print('saved', docx_path)


if __name__ == '__main__':
    md_to_docx(sys.argv[1], sys.argv[2])
