# -*- coding: utf-8 -*-
"""生成 Supplementary File 3（多工作表）与版本溯源表"""
import os, re, ast, json
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = r"d:\Desktop\v8 for JD"
ANA = os.path.join(ROOT, "03_数据", "08_分析用")
LOCK = os.path.join(ROOT, "03_数据", "06_锁定数据集")
FINAL = os.path.join(ROOT, "03_数据", "04_最终表", "最终表.xlsx")
SWT = os.path.join(ROOT, "03_数据", "09_软件表", "软件类别与来源表.csv")
OUTD = os.path.join(ROOT, "02_图表附件", "R2_补充材料")
os.makedirs(OUTD, exist_ok=True)

def nt(s): return re.sub(r"[^a-z0-9]+", "", str(s).lower())
def pl(x):
    try: return ast.literal_eval(x) if isinstance(x, str) else (x if isinstance(x, list) else [])
    except Exception: return []

d = pd.read_csv(os.path.join(ANA, "分析数据集_final.csv"), encoding="utf-8-sig")
old = pd.read_excel(FINAL); old["_n"] = old["Title"].map(nt)
ref = dict(zip(old["_n"], old["Refrence"])); pur = dict(zip(old["_n"], old["Main Purposes"]))
auth = dict(zip(old["_n"], old["Author"]))
d["_soft"] = d["_soft"].map(pl); d["_spec"] = d["_spec"].map(pl); d["_scen"] = d["_scen"].map(pl)
d["Reference"] = d["Title"].map(nt).map(ref)
d["Main purposes"] = d["Title"].map(nt).map(pur)
d["First author"] = d["Title"].map(nt).map(auth)

sw = pd.read_csv(SWT, encoding="utf-8-sig")
cat = dict(zip(sw["规范名称"], sw["类别"])); dom = dict(zip(sw["规范名称"], sw["原始开发领域"]))

studies = pd.DataFrame({
    "Reference": d["Reference"], "First author": d["First author"], "Title": d["Title"], "DOI": d["DOI"],
    "Journal": d["Journal"], "Year": d["Year"], "Region": d["Region"],
    "Study design": d["study_type"].map({"computational": "Computational/simulation", "in_vitro": "In vitro/laboratory",
                                          "clinical": "Clinical study", "case_report": "Case report",
                                          "technical_note": "Technical note", "educational": "Educational"}),
    "Dental specialty": d["Dental Specialty"],
    "Application scenario": d["Application Scenario"],
    "Main purposes": d["Main purposes"],
    "Nondental 3D software used": d["Software Used (fixed)"],
    "Software category(ies)": d["_soft"].map(lambda l: ";".join(sorted({str(cat.get(s, "TO VERIFY")) for s in l}))),
})

glossary = sw[["规范名称", "类别", "原始开发领域", "开发商", "来源URL", "研究数"]].rename(columns={
    "规范名称": "Software", "类别": "Category", "原始开发领域": "Original development domain",
    "开发商": "Developer", "来源URL": "Source (official product page)", "研究数": "Studies (n)"})
CATNAME = {"MIP": "Medical image processing", "GEN3D": "General-purpose 3D modelling/rendering",
           "CAD": "Engineering CAD", "SIM": "Engineering simulation (FEA/CFD)",
           "RE": "Reverse engineering & metrology", "AM": "Additive manufacturing / slicing",
           "PHOTO": "Photogrammetry / 3D scanning", "RT": "Real-time / XR engine",
           "SCI": "Scientific imaging & analysis", "DICOM": "DICOM viewer", "DENTAL": "Dental-specific (not counted)",
           "???": "To verify", "?": "To verify"}
glossary["Category"] = glossary["Category"].map(lambda c: CATNAME.get(c, c))

audit = pd.read_csv(os.path.join(LOCK, "剔除清单_共41条.csv"), encoding="utf-8-sig")
extra = pd.DataFrame([
    {"序号": "", "Title": "Accuracy of dynamic computer-assisted implant surgery in fully edentulous patients",
     "Item Type": "journalArticle", "Date": "", "裁定理由": "PDF verified: only dental-specific software (Navident/EvaluNav)", "审核轮次": "第三轮(PDF核验)"},
    {"序号": "", "Title": "A Novel Virtual Planned-Orthodontic-Surgical Approach for Proportional Condylectomy",
     "Item Type": "journalArticle", "Date": "", "裁定理由": "PDF verified: only dental-specific software (NEMOfab/Nemotec)", "审核轮次": "第三轮(PDF核验)"},
    {"序号": "", "Title": "Modelling growth curves of the normal infant's mandible: 3D measurements",
     "Item Type": "journalArticle", "Date": "", "裁定理由": "No full text available; software (Robins 3D) could not be verified", "审核轮次": "第三轮(PDF核验)"},
])
audit_all = pd.concat([audit[["序号", "Title", "Item Type", "Date", "裁定理由", "审核轮次"]], extra], ignore_index=True)

notes = pd.DataFrame({"Item": [
    "Inclusion criteria", "Exclusion criteria", "Multi-label coding",
    "Screening corpus", "AI-assisted screening", "Full-text verification", "Version"],
 "Definition/Value": [
    "Peer-reviewed original research, technical notes or case reports; any dental discipline; explicit use of nondental 3D software; English; published 2020-07-01 to 2026-06-30.",
    "Reviews (narrative/systematic/scoping/meta-analysis), editorials, commentaries, conference abstracts, preprints, video articles, non-dental topics; studies using only dental-specific software; studies without 3D data processing.",
    "A study may be assigned to more than one dental speciality and to more than one application scenario; all counts of speciality/scenario are therefore reported as study–speciality assignments alongside unique study counts.",
    "2,540 records after de-duplication and removal of retracted (n=4) and non-journal (n=10) records.",
    "deepseek-chat API alias (resolved to deepseek-flash, DeepSeek-V4.1-Flash); temperature 0.1; max_tokens 500; three independent runs; Fleiss' kappa = 0.936. Prompt and complete logs deposited in the public repository.",
    "Software names were checked by reproducible text matching against available PDFs; 455/516 studies (88.2%) with full text had all annotated software names found in the extracted PDF text. This does not verify precise software role or purpose. Geomagic product names were re-coded from full texts (n=128 corrections).",
    "Dataset v2.3 (locked); N = 569. The archived human-verification log for historical AI exclusions was incomplete.",
 ]})

xlsx = os.path.join(OUTD, "Supplementary_File_3.xlsx")
with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
    studies.to_excel(w, sheet_name="1_Included studies", index=False)
    glossary.to_excel(w, sheet_name="2_Software glossary", index=False)
    audit_all.to_excel(w, sheet_name="3_Eligibility audit", index=False)
    notes.to_excel(w, sheet_name="4_Definitions and notes", index=False)
    wb = w.book
    for sname in wb.sheetnames:
        ws = wb[sname]
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for i, col in enumerate(studies.columns if sname.startswith("1") else [], 1):
            ws.column_dimensions[get_column_letter(i)].width = {"Title": 60, "DOI": 30, "Journal": 22,
                "Dental specialty": 26, "Application scenario": 34, "Main purposes": 40,
                "Nondental 3D software used": 26, "Software category(ies)": 30, "Reference": 20}.get(col, 16)
print("已生成:", xlsx)
print("studies:", studies.shape, "| glossary:", glossary.shape, "| audit:", audit_all.shape)

# 版本溯源表
prov = pd.DataFrame([
    {"Output": "Figure 1 (PRISMA)", "Dataset version": "v2.3 (N=569) + search/clean logs", "Notes": "3,631 identified; 1,091 removed; 2,540 screened; full-text files retrieved or available for verification for 1,747 records; 569 included after available human screening records and eligibility audits"},
    {"Output": "Figure 2a (trend)", "Dataset version": "v2.3", "Notes": "n=562 records with month-level date; half-year, 2026-H1 incomplete"},
    {"Output": "Figure 2b (geography)", "Dataset version": "v2.3", "Notes": "51 countries/regions"},
    {"Output": "Figure 2c (journals)", "Dataset version": "v2.3", "Notes": "denominator = 569; 184 journals; 19 records lack journal metadata"},
    {"Output": "Figure 2d (speciality x half-year)", "Dataset version": "v2.3", "Notes": "study–speciality assignments; unique studies annotated"},
    {"Output": "Figure 3 (speciality x scenario)", "Dataset version": "v2.3", "Notes": "assignments; unique study counts annotated"},
    {"Output": "Figure 5 (contingency)", "Dataset version": "v2.3", "Notes": "chi-square on 10 software x 12 specialties; Cramér's V=0.20"},
    {"Output": "Supplementary File 3", "Dataset version": "v2.3", "Notes": "studies, software glossary, eligibility audit, definitions"},
    {"Output": "Supplementary File 2 (AI logs)", "Dataset version": "screening corpus (2,540)", "Notes": "three runs, raw JSON logs, prompt, model metadata"},
    {"Output": "Eligibility audit", "Dataset version": "v2.1→v2.2→v2.3", "Notes": "41 removed in audits 1-2; 2 removed after PDF verification; 1 removed because software could not be verified without full text = 613→569; historical human full-text log incomplete"},
])
prov.to_csv(os.path.join(OUTD, "Dataset_provenance_table.csv"), index=False, encoding="utf-8-sig")
print("版本溯源表:", os.path.join(OUTD, "Dataset_provenance_table.csv"))
