# -*- coding: utf-8 -*-
"""Convert the Chinese manuscript markdown into a Word file with Chinese typography.

Body: 宋体 for Chinese, Times New Roman for Latin and numbers, 12 pt, 1.5 line spacing,
first-line indent of two characters. Headings are bold; the title is centred.
"""
import io, os, re, sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding='utf-8')
ROOT = r'd:\Desktop\v8 for JD'
SRC = os.path.join(ROOT, '01_投稿文件', '中文版', 'Revised_manuscript_R2_中文.md')
DST = os.path.join(ROOT, '01_投稿文件', '中文版', '非牙科3D软件牙科应用_范围综述_中文稿.docx')

LATIN = 'Times New Roman'
CJK = '宋体'

doc = Document()
st = doc.styles['Normal']
st.font.name = LATIN
st.font.size = Pt(12)
rpr = st.element.get_or_add_rPr()
rf = rpr.find(qn('w:rFonts'))
if rf is None:
    rf = OxmlElement('w:rFonts')
    rpr.append(rf)
rf.set(qn('w:ascii'), LATIN)
rf.set(qn('w:hAnsi'), LATIN)
rf.set(qn('w:eastAsia'), CJK)
pf = st.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.5
pf.space_before = Pt(0)
pf.space_after = Pt(6)


def add(text, style=None, indent=False, center=False, bold=False, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = LATIN
    run.font.size = Pt(size or 12)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    run.bold = bold
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)   # two characters at 12 pt
    return p


blocks = io.open(SRC, encoding='utf-8').read().split('\n\n')
for i, b in enumerate(blocks):
    b = b.strip()
    if not b:
        continue
    m = re.match(r'^(#+)\s*(.*)$', b, re.S)
    if m:
        level, head = len(m.group(1)), m.group(2).strip()
        if level == 1:
            add(head, center=True, bold=True, size=16)
        else:
            add(head, bold=True, size=14 if level == 2 else 12)
        continue
    if re.match(r'^\[\d+\] ', b):                     # reference entry
        p = add(b)
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        continue
    add(b, indent=True)

doc.save(DST)
print('saved', DST)
print('paragraphs:', len(doc.paragraphs))
