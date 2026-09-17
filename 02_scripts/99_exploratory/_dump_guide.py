# -*- coding: utf-8 -*-
"""Dump the journal's guide for authors to plain text."""
import os
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(ROOT)
SRC = os.path.join('01_投稿文件', 'JD guide for authors.docx')
OUT = os.path.join('04_代码', '07_文档', '_guide_text.txt')

doc = Document(SRC)
lines = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for tbl in doc.tables:
    lines.append('[TABLE]')
    for row in tbl.rows:
        lines.append(' | '.join(c.text.strip() for c in row.cells))
open(OUT, 'w', encoding='utf-8').write('\n'.join(lines))
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
print('written:', OUT, len(lines), 'lines')
