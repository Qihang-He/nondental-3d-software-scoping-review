# -*- coding: utf-8 -*-
"""List sheet names and first row of each supplementary workbook."""
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
D = r'd:\Desktop\v8 for JD\01_投稿文件\R2_Submission'
for f in ('Supplementary_File_2.xlsx', 'Supplementary_File_3.xlsx', 'Supplementary_File_4.xlsx'):
    wb = openpyxl.load_workbook(D + '\\' + f, read_only=True, data_only=True)
    print('####', f)
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if i >= 0:
                break
        hdr = ', '.join(str(c)[:22] for c in (rows[0] or []) if c is not None)
        print(f'   - {ws.title}  rows={ws.max_row}  cols={ws.max_column}')
        print(f'     header: {hdr[:170]}')
    print()
